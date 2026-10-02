"""Pull the official USCL data (read-only) for scripts/uscl-reconcile.mts.

    U=<franchise email> P=<password> python scripts/uscl-fetch.py <out.json>

Needs: pip install playwright openpyxl && system Chrome. Credentials come from
the environment and are never written anywhere. Nothing on the USCL site is
clicked except sign-in and the My Squad tab (no wishlist / no changes).

Writes {fetched_at, pool:[...export rows...], squad:[{name,category,price,type}],
public:{TEAM NAME:[names...]}}
"""
import io, json, os, re, sys, datetime
import openpyxl
from playwright.sync_api import sync_playwright

BASE = "https://www.urbansportsstudio.in"
out = sys.argv[1]

with sync_playwright() as p:
    b = p.chromium.launch(channel="chrome")
    ctx = b.new_context(viewport={"width": 1440, "height": 1000})
    pg = ctx.new_page()

    # 1. public franchises page: every team block, names listed under it
    pg.goto(BASE + "/franchises", wait_until="networkidle")
    pg.wait_for_timeout(1500)
    text = pg.inner_text("body")
    public = {}
    teams = re.findall(r"\n([A-Z&' ]{6,})\n\n“", text)
    for i, team in enumerate(teams):
        start = text.index("\n" + team + "\n")
        end = text.index("\n" + teams[i + 1] + "\n", start) if i + 1 < len(teams) else len(text)
        block = text[start:end]
        names = [n.strip(" ·") for n in re.findall(r"\n([A-Z][A-Za-z.() ]+?)·?(?=\n)", block)]
        public[team.strip()] = [n for n in names if n and n.upper() != n and not n.startswith("“")]

    # 2. franchise portal: pool export + My Squad
    pg.goto(BASE + "/franchise/login")
    pg.wait_for_timeout(2500)
    pg.fill("input[name=email]", os.environ["U"])
    pg.fill("input[name=password]", os.environ["P"])
    pg.click("button:has-text('SIGN IN')")
    pg.wait_for_url("**/franchise**", timeout=60000)
    pg.wait_for_timeout(6000)

    r = ctx.request.get(BASE + "/franchise/pool-export")
    if r.status != 200:
        raise SystemExit(f"pool export failed: {r.status}")
    ws = openpyxl.load_workbook(io.BytesIO(r.body()), data_only=True).active
    hdr = [c.value for c in ws[1]]
    pool = [dict(zip(hdr, row)) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]

    pg.locator("button, [role=tab]", has_text="My Squad").first.click()
    pg.wait_for_timeout(2500)
    squad = []
    for tr in pg.locator("tr").all():
        cells = [c.strip() for c in tr.inner_text().split("\t") if c.strip()]
        if len(cells) >= 4 and cells[2].startswith("₹"):
            squad.append({"name": cells[0], "category": cells[1], "price": int(re.sub(r"[^\d]", "", cells[2])), "type": cells[3]})
    b.close()

json.dump({"fetched_at": datetime.datetime.now().isoformat(timespec="seconds"), "pool": pool, "squad": squad, "public": public},
          open(out, "w"), indent=1, ensure_ascii=False)
print(f"pool {len(pool)} · our squad {len(squad)} · public teams {len(public)}")
